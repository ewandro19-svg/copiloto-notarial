"""
Gera o template Word (minuta_compra_venda.docx) v2, agora com LOGICA CONDICIONAL
via Jinja2 (docxtpl), cobrindo as variacoes mais comuns de uma compra e venda:

  - Multiplas partes (vendedores e compradores em lista, nao par fixo)
  - Outorga do conjuge (aparece so se o regime de bens exigir)
  - Forma de pagamento: a vista OU financiado (blocos de clausula diferentes)
  - Usufruto reservado (opcional)

IMPORTANTE: os textos abaixo sao GENERICOS/PLACEHOLDER, escritos com base em
praxe notarial comum -- NAO sao a redacao aprovada do seu tabelionato. Troque
cada trecho marcado com [SUBSTITUIR] pelo texto real assim que tiver em maos.
A logica condicional (quem aparece quando) essa sim ja fica pronta pra usar.
"""
from docx import Document
from docx.shared import Pt, Cm
from docx.enum.text import WD_ALIGN_PARAGRAPH

doc = Document()

sec = doc.sections[0]
sec.top_margin = Cm(3)
sec.bottom_margin = Cm(2)
sec.left_margin = Cm(3)
sec.right_margin = Cm(2)

style = doc.styles['Normal']
style.font.name = 'Times New Roman'
style.font.size = Pt(12)


def p(texto, justificar=True, negrito=False, centralizar=False):
    """Cada chamada = 1 paragrafo novo. Tags Jinja ({% %}) SEMPRE em paragrafo
    proprio, separado do texto -- e assim que se evita a tag quebrar no meio
    de runs do Word e o docxtpl nao conseguir encontrar."""
    par = doc.add_paragraph()
    if centralizar:
        par.alignment = WD_ALIGN_PARAGRAPH.CENTER
    elif justificar:
        par.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    run = par.add_run(texto)
    run.bold = negrito
    return par


# ---------------------------------------------------------------------------
p("ESCRITURA PÚBLICA DE COMPRA E VENDA", centralizar=True, negrito=True)
p("")

p(
    "SAIBAM todos quantos esta escritura pública virem que, aos {{ data_ato_extenso }}, "
    "nesta cidade de {{ cartorio.municipio }}, {{ cartorio.uf }}, nesta serventia, "
    "{{ cartorio.nome }}, situada à {{ cartorio.endereco }}, perante mim, "
    "{{ escrevente.nome }}, {{ escrevente.cargo }}, compareceram as partes entre si "
    "justas e contratadas, a saber:"
)
p("")

# --- VENDEDORES (lista -- funciona pra 1 ou vários) -------------------------
p("{% for v in vendedores %}", justificar=False)
p(
    "OUTORGANTE VENDEDOR(A){% if not loop.last or vendedores|length > 1 %} "
    "({{ loop.index }}º){% endif %}: {{ v.nome }}, {{ v.nacionalidade }}, "
    "{{ v.estado_civil }}, {{ v.profissao }}, portador(a) da Cédula de Identidade "
    "RG nº {{ v.rg }}, inscrito(a) no CPF/MF sob o nº {{ v.cpf }}, residente e "
    "domiciliado(a) em {{ v.endereco }}{% if v.conjuge %}, casado(a) com "
    "{{ v.conjuge.nome }}, {{ v.conjuge.nacionalidade }}, {{ v.conjuge.profissao }}, "
    "portador(a) da Cédula de Identidade RG nº {{ v.conjuge.rg }}, inscrito(a) no "
    "CPF/MF sob o nº {{ v.conjuge.cpf }}{% endif %};"
)
p("{% endfor %}", justificar=False)
p("")

# --- OUTORGA UXÓRIA/MARITAL (so aparece se algum vendedor precisar) --------
p(
    "{% if vendedores|selectattr('exige_outorga')|list %}"
    "[SUBSTITUIR pelo texto real de outorga uxória/marital do seu tabelionato: "
    "explicitar que o(a) cônjuge comparece e anui expressamente com a venda, "
    "conforme exige o regime de bens declarado.]"
    "{% endif %}",
    justificar=False,
)
p("")

# --- COMPRADORES (lista) ----------------------------------------------------
p("{% for c in compradores %}", justificar=False)
p(
    "OUTORGANTE COMPRADOR(A){% if compradores|length > 1 %} "
    "({{ loop.index }}º){% endif %}: {{ c.nome }}, {{ c.nacionalidade }}, "
    "{{ c.estado_civil }}, {{ c.profissao }}, portador(a) da Cédula de Identidade "
    "RG nº {{ c.rg }}, inscrito(a) no CPF/MF sob o nº {{ c.cpf }}, residente e "
    "domiciliado(a) em {{ c.endereco }};"
)
p("{% endfor %}", justificar=False)
p("")

p(
    "Os presentes, juridicamente capazes, identificados na forma acima, do que dou "
    "fé, pelo(s) vendedor(es) me foi dito que é(são) legítimo(s) proprietário(s) do "
    "seguinte imóvel:"
)
p("")

p(
    "IMÓVEL: {{ imovel.descricao }}, situado à {{ imovel.endereco }}, matriculado "
    "sob o nº {{ imovel.matricula }} no {{ imovel.oficio_registro }}, com área de "
    "{{ imovel.area }}, contribuinte municipal nº {{ imovel.inscricao_iptu }}."
)
p("")

p(
    "TÍTULO AQUISITIVO: o imóvel foi adquirido pelo(s) vendedor(es) conforme "
    "{{ imovel.titulo_aquisitivo }}, devidamente registrado na matrícula supra."
)
p("")

# --- USUFRUTO (opcional) ----------------------------------------------------
p(
    "{% if usufruto %}"
    "USUFRUTO: fica reservado, em favor de {{ usufruto.beneficiario }}, o usufruto "
    "vitalício sobre o imóvel ora transacionado, nos termos do artigo 1.390 e "
    "seguintes do Código Civil. [SUBSTITUIR pelo texto real de reserva de usufruto "
    "do seu tabelionato.]"
    "{% endif %}",
    justificar=False,
)
p("")

# --- FORMA DE PAGAMENTO: A VISTA vs FINANCIADO ------------------------------
p(
    "{% if negocio.forma_pagamento == 'a_vista' %}"
    "Que, pela presente escritura e nos melhores termos de direito, o(s) "
    "vendedor(es) vende(m), como de fato vendido(s) tem(êm), ao(s) comprador(es), "
    "o imóvel acima descrito e caracterizado, pelo preço certo e ajustado de "
    "{{ negocio.valor_extenso }} ({{ negocio.valor_numerico }}), que confessa(m) e "
    "declara(m) haver recebido, em moeda corrente nacional, dando ao(s) "
    "comprador(es), por este ato, plena, geral e irrevogável quitação, para nada "
    "mais reclamar em tempo algum."
    "{% elif negocio.forma_pagamento == 'financiado' %}"
    "Que, pela presente escritura e nos melhores termos de direito, o(s) "
    "vendedor(es) vende(m), como de fato vendido(s) tem(êm), ao(s) comprador(es), "
    "o imóvel acima descrito e caracterizado, pelo preço certo e ajustado de "
    "{{ negocio.valor_extenso }} ({{ negocio.valor_numerico }}), sendo "
    "{{ negocio.valor_entrada }} pagos neste ato, em moeda corrente nacional, e o "
    "saldo remanescente de {{ negocio.valor_financiado }} financiado junto a "
    "{{ negocio.financiamento.instituicao }}, nos termos do contrato de "
    "financiamento nº {{ negocio.financiamento.contrato_numero }}, que fica "
    "arquivado nesta serventia. [SUBSTITUIR pelo texto real de cláusula de "
    "alienação fiduciária/hipoteca do seu tabelionato.]"
    "{% endif %}"
)
p("")

p(
    "CERTIDÕES APRESENTADAS: foram apresentadas nesta data as seguintes "
    "certidões, válidas e dentro do prazo legal: {{ certidoes.lista }}."
)
p("")

p(
    "IMPOSTO DE TRANSMISSÃO: o ITBI incidente sobre esta transação, no valor de "
    "{{ negocio.itbi_valor }}, foi recolhido conforme guia nº "
    "{{ negocio.itbi_guia }}, que fica arquivada nesta serventia."
)
p("")

p(
    "Assim o disseram, pediram-me e eu lavrei a presente escritura, que, feita e "
    "lhes sendo lida, acharam-na conforme, outorgam e assinam."
)
p("")
p("")

p("_" * 50, centralizar=True, justificar=False)
p("{% for v in vendedores %}{{ v.nome }}{% if not loop.last %} / {% endif %}"
  "{% endfor %} — Outorgante(s) Vendedor(es)", centralizar=True, justificar=False)
p("")
p("_" * 50, centralizar=True, justificar=False)
p("{% for c in compradores %}{{ c.nome }}{% if not loop.last %} / {% endif %}"
  "{% endfor %} — Outorgante(s) Comprador(es)", centralizar=True, justificar=False)

doc.save("templates/minuta_compra_venda.docx")
print("Template v2 (com lógica condicional) criado em templates/minuta_compra_venda.docx")
